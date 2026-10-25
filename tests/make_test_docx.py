"""生成一份结构完整的测试用 Word 文档，用于验证 WordParser + TextChunker

覆盖点：
  1. 四级标题层级（Heading 1-4），检验章节路径栈
  2. 小表 / 超长表（需按行窗口切分）/ 表格单元格内插图
  3. 正文插图、文末插图（检验图片归属与末尾兜底）
  4. 超长段落（>max_chunk_size，触发句子级、字符级降级切分）
  5. 无结尾标点的段落、列表、中英文混排

用法：python tests/make_test_docx.py
产出：test_file/rag_test_doc.docx（以及临时插图 test_file/_fig*.png）
"""
import os
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DOCX = os.path.join(ROOT, "test_file", "rag_test_doc.docx")


def make_fig(path, label, color):
  """生成示例插图（纯色底 + 英文标签，避免中文字体缺失导致乱码）"""
  img = Image.new("RGB", (560, 320), color)
  d = ImageDraw.Draw(img)
  d.rectangle([20, 20, 540, 300], outline="white", width=3)
  d.text((40, 140), label, fill="white")
  img.save(path)
  return path


def build(text_only_paragraph):
  doc = Document()
  doc.add_heading("RAG 文档处理管线测试报告", level=0)
  doc.add_paragraph("本文件由脚本自动生成，用于验证 Word 解析与分块的完整链路，"
                    "覆盖标题层级、表格切分、图片归属和超长段落降级等场景。")

  # ---------- 第一章 ----------
  doc.add_heading("第一章 文档解析层", level=1)
  doc.add_paragraph("文档解析层负责把 PDF、Word 等原始文件转换成统一的 markdown 表示。"
                    "不同格式的解析难点不同：PDF 依赖版式模型判断标题、段落和表格区域，"
                    "Word 则直接读取 OOXML 结构标签，几乎不存在识别误差。")

  doc.add_heading("1.1 PDF 解析", level=2)
  doc.add_paragraph("PDF 没有语义结构，全部是绝对坐标的文字块。解析器需要用版式模型把"
                    "文字块重新组织成阅读顺序，再用表格模型还原表格的行列关系。"
                    "这一步的误差主要来自跨页表格和图片中的文字。")
  doc.add_paragraph().add_run().add_picture(make_fig(os.path.join(ROOT, "test_file", "_fig1.png"),
                                                     "Figure 1: PDF parsing pipeline",
                                                     (32, 84, 147)), width=Inches(4.2))
  doc.add_paragraph("图 1 展示了 PDF 解析的完整流程：版式识别、表格还原、图片抽取三个环节串行执行。")

  doc.add_heading("1.2 Word 解析", level=2)
  doc.add_paragraph("docx 本质是一个 zip 包，内部用 XML 描述段落、表格、图片和样式。"
                    "python-docx 直接遍历 XML 节点，因此读出的文字与原文完全一致，"
                    "不需要任何模型推理，也没有 OCR 噪音。")
  doc.add_heading("1.2.1 读取顺序", level=3)
  doc.add_paragraph("paragraphs 和 tables 两个集合会丢失段落与表格的真实交错顺序，"
                    "必须遍历 body 的子节点才能按原序输出。")
  doc.add_heading("1.2.2 图片提取", level=3)
  doc.add_paragraph("图片通过关系 id 指向包内的 media 目录。新格式用 a:blip，"
                    "老 .doc 转换来的文件用 v:imagedata，两种标签都要兼容。")

  # ---------- 第二章 ----------
  doc.add_heading("第二章 分块策略", level=1)
  doc.add_paragraph("分块的目标是在不丢失内容的前提下，让每一块语义完整、长度可控。"
                    "当前策略采用逐级降级：优先按段落切，段落超长再按句子切，"
                    "单句仍超长才退到按字符切。")

  doc.add_heading("2.1 长段落处理", level=2)
  doc.add_paragraph(text_only_paragraph)
  doc.add_paragraph("上面这段是刻意构造的超长段落，用于验证句子级切分是否保留重叠缓冲。")

  doc.add_heading("2.2 策略对比", level=2)
  doc.add_paragraph("表 1 列出了三种基础策略的适用场景与代价。")
  small = doc.add_table(rows=1, cols=4)
  small.style = "Table Grid"
  for i, h in enumerate(["策略", "切分依据", "优点", "缺点"]):
    small.rows[0].cells[i].text = h
  rows = [
    ("段落", "空行", "语义完整", "长段落不可控"),
    ("句子", "句末标点", "长度均匀", "跨句上下文丢失"),
    ("字符", "固定长度", "严格可控", "可能切断词语"),
  ]
  for r in rows:
    cells = small.add_row().cells
    for i, v in enumerate(r):
      cells[i].text = v
  doc.add_paragraph("实际工程中通常把三者串成一条降级链，而不是只选其中一种。")

  # ---------- 第三章 ----------
  doc.add_heading("第三章 量化评估", level=1)
  doc.add_heading("3.1 评分维度", level=2)
  doc.add_paragraph("评估采用百分制，五个维度分别衡量完整性、块大小合规、表格完整性、"
                    "粒度合理性和结构可追溯性。表 2 是逐块评分的明细，"
                    "行数较多，会超过单块上限，用于验证表格按行窗口切分并重复表头。")
  big = doc.add_table(rows=1, cols=6)
  big.style = "Table Grid"
  heads = ["编号", "维度", "权重", "实测值", "得分", "口径说明"]
  for i, h in enumerate(heads):
    big.rows[0].cells[i].text = h
  for n in range(1, 34):
    cells = big.add_row().cells
    cells[0].text = f"R{n:02d}"
    cells[1].text = f"第{n}项检查点"
    cells[2].text = f"{2 + n % 7} 分"
    cells[3].text = f"实测 {n * 3} 字"
    cells[4].text = f"{1 + n % 5}.5 分"
    cells[5].text = f"按第 {n} 条口径逐块核对，未发现异常"
  doc.add_paragraph("表格切分后每个子块都要重复表头，并标注行号范围，"
                    "否则检索命中子块时无法判断字段含义。")

  doc.add_heading("3.2 图表混排", level=2)
  doc.add_paragraph("为了验证图片与正文的归属关系，下面插入两张正文插图。")
  doc.add_paragraph().add_run().add_picture(make_fig(os.path.join(ROOT, "test_file", "_fig2.png"),
                                                     "Figure 2: chunk scoring",
                                                     (16, 122, 87)), width=Inches(4.2))
  doc.add_paragraph("图 2 为各维度得分分布。")
  doc.add_paragraph().add_run().add_picture(make_fig(os.path.join(ROOT, "test_file", "_fig3.png"),
                                                     "Figure 3: chunk length distribution",
                                                     (133, 63, 34)), width=Inches(4.2))
  doc.add_paragraph("图 3 为块长度分布，绝大多数块落在理想区间内。")

  doc.add_heading("3.2.1 表格内插图", level=3)
  doc.add_heading("3.2.1.1 单元格嵌套测试", level=4)
  doc.add_paragraph("部分实验报告的截图直接粘贴在表格单元格中，解析后图片引用行会排在表格后面，"
                    "需要归属到上一个已输出的表格块。")
  cell_table = doc.add_table(rows=2, cols=2)
  cell_table.style = "Table Grid"
  cell_table.rows[0].cells[0].text = "检查项"
  cell_table.rows[0].cells[1].text = "截图"
  cell_table.rows[1].cells[0].text = "表格内图片归属"
  cell_table.rows[1].cells[1].paragraphs[0].add_run().add_picture(
    make_fig(os.path.join(ROOT, "test_file", "_fig4.png"), "Figure 4: image in table cell",
             (68, 68, 68)), width=Inches(2.4))
  doc.add_paragraph("表 3 的右列内嵌一张截图，用于验证单元格图片提取。")

  doc.add_heading("第四章 结论", level=1)
  doc.add_paragraph("测试文档覆盖了标题、表格、图片、长段落四类结构。")
  for item in ["标题层级 1-4 全部命中", "超长表格按行窗口切分且保留表头",
               "正文图片与表格图片均正确归属", "超长段落降级切分无内容丢失"]:
    doc.add_paragraph(item, style="List Bullet")
  doc.add_paragraph("以上结果由 tests/test_word.py 自动统计。")

  # 文末遗留图片：后面没有正文，验证兜底挂到最后一块
  doc.add_paragraph().add_run().add_picture(make_fig(os.path.join(ROOT, "test_file", "_fig5.png"),
                                                     "Figure 5: trailing image",
                                                     (94, 54, 132)), width=Inches(3.0))
  return doc


def long_paragraph():
  """构造一段超过 1000 字的正文，模拟真实文档里的大段论述，末尾故意不带标点"""
  base = [
    "在真实业务文档中，经常会出现一整段连续论述没有任何空行分隔的情况，这类段落往往包含多个并列的观点，",
    "每个观点内部又带有补充说明和举例，读起来内容连贯，但对分块器来说却是一段无法按段落切开的整体。",
    "如果直接按最大块长度硬切，很容易在句子中间断开，把一个完整的技术描述拆成两半，",
    "检索时命中的半句往往语义残缺，反而降低了召回内容的可用性。",
    "因此需要先按句号、问号、分号等句末标点切出句子，再按长度把句子重新组装成块，",
    "这样每一块都由若干个完整句子组成，语义边界与原文一致。",
    "当某一句本身就超过最大块长度时，说明它内部还有可以细分的位置，此时才退到按字符切分，",
    "并且只在切分点附近保留一小段重叠，让相邻块共享上下文，避免关键词恰好被切在边界上。",
    "需要强调的是，重叠只应该出现在同一段的内部切分场景，跨段落之间不做重叠，",
    "因为相邻两段讨论的主题可能完全不同，重叠只会把无关内容带进检索结果。",
    "最后还有一种边界情况：段落末尾的最后一句话没有句末标点，例如列举到一半就结束的短句，",
    "此时切分逻辑不能把它当作无效内容丢弃，而应该并入上一块。",
    "在实际的课程实验报告、技术白皮书和产品需求文档中，这种长段落非常常见，",
    "尤其是需求背景、方案论证和原理阐述这几类章节，作者往往习惯把多个理由写在同一段里，",
    "用逗号、顿号层层递进，整段读下来一气呵成，但字数轻松就超过一千。",
    "如果分块器只按空行切分，这样的段落就会变成一个超大块，超过模型上下文限制后被截断，",
    "检索时也只能命中整段，精度很差。因此在段落内部继续降级是必要的工程手段。",
    "降级链的顺序也很讲究，必须先句子后字符，因为句子的语义边界是作者明确写出来的，",
    "而字符边界完全是机械截断，能不用就不用。句子级切分时还要注意保留重叠，",
    "让相邻块在边界处共享一小段上下文，这样即使某个关键词恰好落在边界上，也不会两边都漏。",
    "重叠长度不能太大，否则块之间的重复内容会迅速膨胀，浪费存储和向量化成本，",
    "一般取最大块长的百分之十到百分之二十比较合适，具体取值还要结合文档的语言特点。",
    "中文的字均信息量高于英文，同样字数下块可以略小一些，重叠也可以略短一些。",
    "此外还要防止一种隐蔽的错误：切分后把所有子块都当成独立块输出，",
    "却忘了给它们标注所属章节，结果检索命中后无法回溯到原文位置，",
    "排查问题时只能重新打开原始文档人工查找，效率极低。",
    "最后还有一种边界情况：段落末尾的最后一句话没有句末标点，例如列举到一半就结束的短句，",
    "此时切分逻辑不能把它当作无效内容丢弃，而应该并入上一块",
  ]
  return "".join(base)


def main():
  doc = build(long_paragraph())
  os.makedirs(os.path.dirname(OUT_DOCX), exist_ok=True)
  doc.save(OUT_DOCX)
  print(f"已生成 {OUT_DOCX}")


if __name__ == "__main__":
  main()
