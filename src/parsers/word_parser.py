import io
import os
import re
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.parsers.base import BaseParser, ParseResult

# 内嵌图片在 run 里的关系 id 命名空间
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _iter_block_items(doc):
    """按文档真实顺序产出段落和表格（doc.paragraphs / doc.tables 会丢失交错顺序）"""
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}p'):
            yield Paragraph(child, doc)
        elif child.tag.endswith('}tbl'):
            yield Table(child, doc)


def _heading_level(style_name: str):
    """标题样式名 → 层级：英文 'Heading 1'、中文 '标题 1'，取不到返回 None"""
    m = re.match(r'^(?:Heading|标题)\s*(\d)', style_name or '')
    return int(m.group(1)) if m else None


def _cell_text(cell) -> str:
    # 单元格内的换行和竖线会破坏 markdown 表格，压平并转义
    return cell.text.replace("\n", " ").replace("|", "\\|").strip()


class WordParser(BaseParser):
  """基于 python-docx 的 Word 解析器：docx 是 zip+XML，直接读结构，不用模型"""

  def __init__(self, image_dir: str = "data/images"):
    self.image_dir = image_dir

  def parse(self, file_path: str) -> ParseResult:
    doc = Document(file_path)
    raw_text, tables = self._to_markdown(doc, Path(file_path).stem)
    return ParseResult(
      filename=Path(file_path).name,
      total_pages=0,   # Word 没有页的概念
      total_chunks=len(raw_text),
      table_count=len(tables),
      chunks=[],
      raw_text=raw_text,
      tables=tables,
      metadata={"parser": "python-docx", "file_path": file_path}
    )

  def parse_from_bytes(self, file_bytes: bytes, filename: str) -> ParseResult:
    doc = Document(io.BytesIO(file_bytes))
    raw_text, tables = self._to_markdown(doc, Path(filename).stem)
    return ParseResult(
      filename=filename,
      total_pages=0,
      total_chunks=len(raw_text),
      table_count=len(tables),
      chunks=[],
      raw_text=raw_text,
      tables=tables,
      metadata={"parser": "python-docx", "file_path": filename}
    )

  def _to_markdown(self, doc, stem: str) -> Tuple[str, List[Dict]]:
    """python-docx 对象 → markdown：标题映射 #，表格拼 | 表，图片存本地留引用"""
    lines = []
    tables = []
    self._img_count = 0
    self._doc = doc
    self._stem = stem

    for block in _iter_block_items(doc):
      if isinstance(block, Table):
        md = self._table_to_markdown(block)
        if not md:
          continue
        tables.append({"table_number": len(tables) + 1, "page": None, "markdown": md})
        lines.append(md)
        # 表格单元格里的图片（实验报告的截图通常都插在表格里）
        for row in block.rows:
          for cell in row.cells:
            lines.extend(self._collect_images(cell.paragraphs))
        lines.append("")
        continue

      text = block.text.strip()
      level = _heading_level(block.style.name if block.style else None)
      if level:
        lines.append(f"{'#' * min(level, 6)} {text}")
      elif text.strip():
        style = block.style.name if block.style else ''
        if style.startswith('List'):
          lines.append(f"- {text.strip()}")
        else:
          lines.append(text)

      lines.extend(self._collect_images([block]))
      lines.append("")

    return "\n".join(lines), tables

  def _collect_images(self, paragraphs) -> List[str]:
    """从一组段落里提取内嵌/浮动图片，存本地并返回 markdown 引用行"""
    refs = []
    for para in paragraphs:
      for run in para.runs:
        # 两种图片来源：a:blip（新格式 DrawingML）和 v:imagedata（老 .doc 转来的 VML），
        # 用 local-name() 查询避免命名空间前缀在不同元素上注册不一致的问题
        nodes = run._element.xpath(".//*[local-name()='blip']") + \
                run._element.xpath(".//*[local-name()='imagedata']")
        for node in nodes:
          rid = node.get(R_EMBED) or node.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
          if not rid or rid not in self._doc.part.related_parts:
            continue
          part = self._doc.part.related_parts[rid]
          self._img_count += 1
          ext = os.path.splitext(str(part.partname))[1] or ".png"
          path = os.path.join(self.image_dir, f"{self._stem}_{self._img_count}{ext}")
          os.makedirs(self.image_dir, exist_ok=True)
          with open(path, "wb") as f:
            f.write(part.blob)
          refs.append(f"![图片{self._img_count}]({path.replace(os.sep, '/')})")
    return refs

  @staticmethod
  def _table_to_markdown(table: Table) -> str:
    rows = [[_cell_text(c) for c in row.cells] for row in table.rows]
    if not rows:
      return ""
    header = "| " + " | ".join(rows[0]) + " |"
    sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([header, sep] + body)

  def supported_extensions(self) -> List[str]:
    return [".docx"]
